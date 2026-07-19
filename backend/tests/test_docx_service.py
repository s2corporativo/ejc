"""Smoke test da geração de peça em DOCX (docx_service.gerar_docx).

O caminho DOCX (peça editável) não tinha nenhum teste — o bump do python-docx
(PR #131) expôs a lacuna. Aqui exercitamos gerar_docx ponta a ponta: markdown
básico → bytes de um .docx VÁLIDO e abrível, com o título e o corpo presentes.
Protege um caminho crítico (exportar peça) contra regressão silenciosa da lib.
"""
from __future__ import annotations

import io

from app.services.docx_service import gerar_docx


CONTEUDO_MD = """# Petição Inicial

Excelentíssimo Senhor Doutor Juiz de Direito.

## DOS FATOS

O autor **sofreu dano** e requer *reparação*.

- primeiro ponto
- segundo ponto

1. pedido um
2. pedido dois

> Citação longa em bloco, no padrão ABNT.
"""


def test_gerar_docx_retorna_docx_valido_e_abrivel():
    dados = gerar_docx(
        "Petição Inicial",
        CONTEUDO_MD,
        meta={"numero_processo": "0001234-56.2026.8.13.0027", "codigo_peca": "EJC-CIV-001"},
    )

    # É bytes não-vazio e tem a assinatura ZIP (todo .docx é um zip "PK\x03\x04").
    assert isinstance(dados, bytes)
    assert len(dados) > 0
    assert dados[:2] == b"PK"

    # Abre de volta com python-docx e confere que título e corpo entraram.
    from docx import Document

    doc = Document(io.BytesIO(dados))
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "PETIÇÃO INICIAL" in texto        # título é uppercased
    assert "DOS FATOS" in texto              # heading nível 2
    assert "sofreu dano" in texto            # negrito preservado como texto
    assert "primeiro ponto" in texto         # item de lista


def test_gerar_docx_sem_meta_e_conteudo_vazio_nao_quebra():
    """Robustez: sem meta e com conteúdo vazio ainda produz um .docx válido
    (só o cabeçalho institucional/título), sem lançar."""
    dados = gerar_docx("Documento", "", meta=None)
    assert isinstance(dados, bytes) and dados[:2] == b"PK"

    from docx import Document

    doc = Document(io.BytesIO(dados))
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "DOCUMENTO" in texto
