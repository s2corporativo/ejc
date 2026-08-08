"""Marca de origem-IA que VIAJA com o documento + dados FIXOS do escritório.

Auditoria de geração de documentos:
  (A/B) OAB, endereço, CEP e CNPJ do escritório vêm da FONTE ÚNICA (settings
        ESCRITORIO_*); vazio no .env OMITE o segmento inteiro do timbre — nunca
        dado inventado, "A PREENCHER" mudo ou placeholder de pendência interna
        no documento entregue ao cliente (a pendência vai para o log de boot).
  (C)   O aviso de rascunho-IA agora é EMBUTIDO no PDF/DOCX exportado quando a
        peça é ai_generated e ainda NÃO foi human_reviewed (minuta não revisada);
        a versão revisada sai LIMPA. Sem gate de bloqueio no download.
"""
from __future__ import annotations

import io

from app.core.config import Settings, get_settings
from app.services.docx_service import gerar_docx
from app.services.pdf_service import _texto_peca_para_html


_MARCA = "MINUTA GERADA POR IA"


# ── (C) DOCX: a marca viaja com o arquivo, condicional à revisão ──────────────
def test_docx_minuta_nao_revisada_contem_marca_ia():
    dados = gerar_docx("Petição Inicial", "# Petição\n\nCorpo.", meta={"minuta_ia": True})
    doc = _abrir_docx(dados)
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert _MARCA in texto
    assert "NAO PROTOCOLAR SEM REVISAO" in texto


def test_docx_peca_revisada_sai_limpa_sem_marca_ia():
    # human_reviewed=True → o router não seta minuta_ia; versão final limpa.
    dados = gerar_docx("Petição Inicial", "# Petição\n\nCorpo.", meta={"minuta_ia": False})
    doc = _abrir_docx(dados)
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert _MARCA not in texto


def test_docx_sem_meta_nao_marca_como_minuta_ia():
    dados = gerar_docx("Documento", "corpo", meta=None)
    doc = _abrir_docx(dados)
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert _MARCA not in texto


# ── (C) PDF (caminho HTML puro, sem depender de weasyprint) ───────────────────
def test_pdf_html_minuta_nao_revisada_contem_marca_ia():
    html = _texto_peca_para_html("Petição", "DOS FATOS\n\nCorpo.", minuta_ia=True)
    assert _MARCA in html
    assert 'class="ia-minuta"' in html


def test_pdf_html_peca_revisada_sai_limpa_sem_marca_ia():
    html = _texto_peca_para_html("Petição", "DOS FATOS\n\nCorpo.", minuta_ia=False)
    assert _MARCA not in html
    assert "ia-minuta" not in html


# ── (A/B) Dados FIXOS do escritório: fonte única, sem placeholder no papel ────
# Onda 1 da refatoração: setting institucional vazia NÃO imprime mais
# "[CEP - preencher em .env]" no timbre — o segmento inteiro é omitido e a
# pendência migra para escritorio_pendencias() (log de boot / diagnóstico).
def test_helper_devolve_vazio_quando_setting_vazia():
    s = Settings(ESCRITORIO_CEP="   ", ESCRITORIO_OAB="")
    assert s.escritorio_cep() == ""
    assert s.escritorio_oab() == ""


def test_helper_normaliza_setting_preenchida():
    assert Settings(ESCRITORIO_OAB="  12.345 ").escritorio_oab() == "12.345"


def test_pendencias_listam_apenas_os_campos_vazios():
    s = Settings(ESCRITORIO_OAB="251174", ESCRITORIO_ENDERECO="Rua X, 10", ESCRITORIO_CEP="")
    # CNPJ nasce vazio (auditoria jul/2026) — entra na pendência junto do CEP.
    assert s.escritorio_pendencias() == ["ESCRITORIO_CEP", "ESCRITORIO_CNPJ"]
    assert Settings(
        ESCRITORIO_CEP="32510-010", ESCRITORIO_CNPJ="11.222.333/0001-81"
    ).escritorio_pendencias() == []


def test_timbre_docx_omite_segmento_vazio_em_vez_de_imprimir_placeholder():
    """Regressão: com CEP vazio (default), nem o placeholder nem o rótulo órfão
    'CEP' podem aparecer no timbre/rodapé do documento entregue ao cliente."""
    from app.services import docx_service

    assert "preencher em .env" not in docx_service.ESCRITORIO_TIMBRE_SUB
    assert "preencher em .env" not in docx_service.ESCRITORIO_CONTATO
    if not docx_service.settings.escritorio_cep():
        assert "CEP" not in docx_service.ESCRITORIO_TIMBRE_SUB
        assert "CEP" not in docx_service.ESCRITORIO_CONTATO
    # Sem rótulo órfão nem separador duplicado sobrando.
    for linha in (docx_service.ESCRITORIO_TIMBRE_SUB, docx_service.ESCRITORIO_CONTATO):
        assert "|  |" not in linha
        assert not linha.strip().endswith("|")


def test_timbre_pdf_omite_segmento_vazio():
    from app.services import pdf_service

    for linha in (pdf_service._ESC_SUB1, pdf_service._ESC_SUB2):
        assert "preencher em .env" not in linha
        assert not linha.strip().endswith("&nbsp;|&nbsp;")


def test_docx_timbre_traz_oab_e_endereco_do_escritorio():
    from app.services import docx_service

    dados = gerar_docx("Peça", "corpo", meta=None)
    doc = _abrir_docx(dados)
    header = doc.sections[0].header
    texto_header = "\n".join(p.text for p in header.paragraphs)
    assert "OAB/MG" in texto_header
    assert docx_service.settings.escritorio_endereco() in texto_header


def test_procuracao_usa_dados_do_escritorio_e_remove_placeholders_legados():
    from datetime import date

    from app.models.case import Case
    from app.models.client import Client
    from app.services.documental import _procuracao

    s = get_settings()
    case = Case(titulo="Cobranca X", numero_processo=None)
    cli = Client(nome="Fulano de Tal")
    # OUTORGADO agora é FIXO (sócio-titular): o `adv` passado NÃO alimenta mais o
    # outorgado — o modelo oficial é sempre outorgado ao sócio.
    texto = _procuracao(case, cli, "Dra. Beltrana")

    # Cabeçalho oficial (timbre textual) + cidade/UF de assinatura das settings.
    assert "DE PAULA TEIXEIRA - SOCIEDADE DE ADVOGADOS" in texto
    assert f"{s.ESCRITORIO_CIDADE}/{s.ESCRITORIO_ESTADO}," in texto
    # OUTORGADO FIXO: sócio-titular com OAB/MG 251.174 — não o `adv` da chamada.
    assert "JOÃO PEDRO RODRIGUES TEIXEIRA" in texto
    assert "251.174" in texto
    assert "Dra. Beltrana" not in texto
    # Placeholders legados de dado FIXO NÃO escapam mais para o documento.
    assert "[OAB/MG no ____]" not in texto
    assert "[endereco do escritorio]" not in texto
    assert "[Cidade]/MG" not in texto
    # A DATA agora é a corrente (por extenso) — o placeholder [data] some.
    assert "[data]" not in texto
    assert str(date.today().year) in texto


# ── (Onda 2) Rótulos internos NÃO vazam para o documento do cliente ──────────
# A auditoria de julho/2026 flagrou a marca interna "Visual Law EJC" na capa do
# PDF e no quadro de controle/nota do DOCX. O quadro de controle continua
# existindo (dados úteis), mas sem a marca; a busca é case-insensitive e
# ignora acentuação para a regressão não escapar por grafia.
def _sem_acentos(texto: str) -> str:
    import unicodedata

    return "".join(
        c for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn"
    )


def _contem_visual_law(texto: str) -> bool:
    return "VISUAL LAW" in _sem_acentos(texto).upper()


def _texto_completo_docx(doc) -> str:
    """Todo texto visível do DOCX: corpo, tabelas, cabeçalho e rodapé."""
    textos = [p.text for p in doc.paragraphs]
    for tabela in doc.tables:
        for cell in tabela._cells:
            textos.extend(p.text for p in cell.paragraphs)
    for section in doc.sections:
        textos.extend(p.text for p in section.header.paragraphs)
        textos.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(textos)


def test_pdf_html_nao_contem_marca_visual_law_e_kicker_e_neutro_acentuado():
    import re as _re

    html = _texto_peca_para_html("Petição Inicial", "DOS FATOS\n\nCorpo.", minuta_ia=True)
    # Só o texto VISÍVEL conta (a classe CSS interna "visual-law" não imprime).
    visivel = _re.sub(r"<[^>]+>", " ", html)
    assert not _contem_visual_law(visivel)
    # Kicker neutro E acentuado (pipeline de acentuação preservado — PR #765).
    assert "Peça jurídica" in html
    assert "Peca juridica" not in html


def test_docx_nao_contem_marca_visual_law_e_mantem_quadro_de_controle():
    dados = gerar_docx(
        "Petição Inicial", "# Petição\n\nCorpo.",
        meta={"minuta_ia": True, "codigo_peca": "EJC-2026-0001", "versao": 2},
    )
    texto = _texto_completo_docx(_abrir_docx(dados))
    assert not _contem_visual_law(texto)
    # O quadro de controle sobrevive — só a marca interna some.
    assert "CONTROLE DO DOCUMENTO" in texto
    assert "EJC-2026-0001" in texto


def test_docx_sem_codigo_peca_usa_fallback_neutro_no_controle():
    dados = gerar_docx("Documento", "corpo", meta=None)
    texto = _texto_completo_docx(_abrir_docx(dados))
    assert not _contem_visual_law(texto)


# ── (Onda 2) CNPJ do timbre: sem default hardcoded, pendência visível ─────────
# Auditoria jul/2026: o CNPJ fixo no código resolvia para OUTRA razão social.
# Mesmo padrão do CEP: default vazio → segmento omitido do timbre + pendência
# em escritorio_pendencias() (log de boot / Central de Diagnóstico).
_CNPJ_ANTIGO = "32.491.468"  # prefixo do valor hardcoded removido


def test_timbre_sem_cnpj_nao_imprime_placeholder_nem_cnpj_antigo():
    from app.services import docx_service, pdf_service

    linhas = (docx_service.ESCRITORIO_CONTATO, pdf_service._ESC_SUB1)
    for linha in linhas:
        assert _CNPJ_ANTIGO not in linha
        assert "preencher" not in linha
    if not docx_service.settings.escritorio_cnpj():
        # Sem CNPJ no .env: nem o rótulo órfão "CNPJ" pode sobrar no timbre.
        for linha in linhas:
            assert "CNPJ" not in linha


def test_rodape_dos_templates_sem_cnpj_nao_tem_rotulo_orfao():
    from app.services.system_prompts import templates_documentos as td

    assert _CNPJ_ANTIGO not in td.RODAPE
    if not td.DADOS_ESCRITORIO["cnpj"]:
        assert "CNPJ" not in td.RODAPE


def test_cnpj_vazio_entra_nas_pendencias_e_preenchido_sai():
    assert "ESCRITORIO_CNPJ" in Settings(ESCRITORIO_CNPJ="  ").escritorio_pendencias()
    assert "ESCRITORIO_CNPJ" not in Settings(
        ESCRITORIO_CNPJ="11.222.333/0001-81"
    ).escritorio_pendencias()


def test_helper_escritorio_cnpj_normaliza_e_devolve_vazio():
    assert Settings(ESCRITORIO_CNPJ="").escritorio_cnpj() == ""
    assert Settings(ESCRITORIO_CNPJ=" 11.222.333/0001-81 ").escritorio_cnpj() == "11.222.333/0001-81"


def _abrir_docx(dados: bytes):
    assert isinstance(dados, bytes) and dados[:2] == b"PK"
    from docx import Document

    return Document(io.BytesIO(dados))
