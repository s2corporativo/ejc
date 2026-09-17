# ── app/services/bank_report.py ──────────────────────────────────────────────
# Gera a planilha Excel (DASHBOARD + EXTRATOS + COBRANÇAS) e os documentos
# jurídicos (Notificação Extrajudicial, Petição Inicial, Reclamação BACEN).
# Tudo é MINUTA — revisão obrigatória do advogado (OAB). Nunca promete resultado.
from __future__ import annotations
import io
import html
from datetime import date


# ─────────────────────────────── Excel ───────────────────────────────────────
def gerar_excel(analise: dict, transacoes: list[dict], cobrancas: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    navy = PatternFill("solid", fgColor="1F2A44")
    red = PatternFill("solid", fgColor="F8D7DA")
    hdrf = Font(bold=True, color="FFFFFF")
    bold = Font(bold=True)
    center = Alignment(horizontal="center")

    def head(ws, cols, larguras=None):
        for i, c in enumerate(cols, 1):
            cell = ws.cell(1, i, c)
            cell.fill = navy; cell.font = hdrf; cell.alignment = center
        if larguras:
            for i, w in enumerate(larguras, 1):
                ws.column_dimensions[ws.cell(1, i).column_letter].width = w

    # ── DASHBOARD ──
    ws = wb.active; ws.title = "DASHBOARD"
    ws.merge_cells("A1:D1")
    t = ws.cell(1, 1, "ANÁLISE BANCÁRIA — RELATÓRIO"); t.fill = navy; t.font = Font(bold=True, color="FFFFFF", size=14); t.alignment = center
    linhas = [
        ("Banco", analise.get("banco") or "—"),
        ("Arquivo", analise.get("arquivo_nome") or "—"),
        ("Período", f"{analise.get('periodo_inicio') or '?'} a {analise.get('periodo_fim') or '?'}"),
        ("Transações", analise.get("total_transacoes", 0)),
        ("Total créditos (R$)", float(analise.get("total_creditos") or 0)),
        ("Total débitos (R$)", float(analise.get("total_debitos") or 0)),
        ("Cobranças abusivas (qtd)", analise.get("qtd_abusivas", 0)),
        ("Total potencialmente indevido (R$)", float(analise.get("total_abusivo") or 0)),
    ]
    for i, (k, v) in enumerate(linhas, 3):
        ws.cell(i, 1, k).font = bold
        ws.cell(i, 2, v)
    ws.cell(12, 1, "Resumo por tipo de cobrança").font = bold
    ws.cell(13, 1, "Regra").font = hdrf; ws.cell(13, 1).fill = navy
    ws.cell(13, 2, "Qtd").font = hdrf; ws.cell(13, 2).fill = navy
    ws.cell(13, 3, "Valor (R$)").font = hdrf; ws.cell(13, 3).fill = navy
    resumo: dict = {}
    for c in cobrancas:
        r = resumo.setdefault(c["regra"], [0, 0.0])
        r[0] += 1; r[1] += float(c.get("valor") or 0)
    for i, (regra, (qtd, val)) in enumerate(sorted(resumo.items(), key=lambda x: -x[1][1]), 14):
        ws.cell(i, 1, regra); ws.cell(i, 2, qtd); ws.cell(i, 3, round(val, 2))
    ws.column_dimensions["A"].width = 38; ws.column_dimensions["B"].width = 18; ws.column_dimensions["C"].width = 18

    # ── EXTRATOS ──
    we = wb.create_sheet("EXTRATOS")
    head(we, ["Data", "Descrição", "Valor (R$)", "Tipo", "Saldo (R$)"], [14, 60, 16, 12, 16])
    for i, t in enumerate(transacoes, 2):
        we.cell(i, 1, t["data"].strftime("%d/%m/%Y") if t.get("data") else "")
        we.cell(i, 2, t.get("descricao") or "")
        we.cell(i, 3, t.get("valor"))
        we.cell(i, 4, t.get("tipo") or "")
        we.cell(i, 5, t.get("saldo"))

    # ── COBRANÇAS ──
    wc = wb.create_sheet("COBRANÇAS")
    head(wc, ["Data", "Descrição", "Regra", "Base legal", "Prioridade", "Valor (R$)"], [14, 44, 20, 34, 14, 16])
    for i, c in enumerate(cobrancas, 2):
        tx = c.get("transaction") or {}
        wc.cell(i, 1, tx.get("data").strftime("%d/%m/%Y") if tx.get("data") else "")
        wc.cell(i, 2, c.get("titulo") or "")
        wc.cell(i, 3, c.get("regra") or "")
        wc.cell(i, 4, c.get("base_legal") or "")
        pr = wc.cell(i, 5, c.get("prioridade") or "")
        if (c.get("prioridade") or "") == "URGENTE":
            pr.fill = red; pr.font = bold
        wc.cell(i, 6, float(c.get("valor") or 0))

    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


# ───────────────────────────── Documentos ────────────────────────────────────
def _linhas_cobrancas_html(cobrancas: list[dict]) -> str:
    linhas = []
    for c in cobrancas:
        tx = c.get("transaction") or {}
        d = tx.get("data").strftime("%d/%m/%Y") if tx.get("data") else "—"
        linhas.append(
            f"<tr><td>{d}</td><td>{html.escape(c.get('titulo') or '')}</td>"
            f"<td>{html.escape(c.get('base_legal') or '')}</td>"
            f"<td style='text-align:right'>R$ {float(c.get('valor') or 0):.2f}</td></tr>"
        )
    return "\n".join(linhas)


def _doc_base(titulo: str, corpo: str) -> str:
    return f"""<!doctype html><html lang="pt-br"><head><meta charset="utf-8">
<title>{html.escape(titulo)}</title>
<style>
body{{font-family:'Times New Roman',serif;max-width:760px;margin:32px auto;color:#1a1a1a;line-height:1.6;font-size:14px}}
h1{{font-size:16px;text-align:center;text-transform:uppercase;letter-spacing:1px}}
table{{width:100%;border-collapse:collapse;margin:14px 0;font-size:12.5px}}
th,td{{border:1px solid #999;padding:5px 7px;text-align:left}}
th{{background:#1F2A44;color:#fff}}
.aviso{{margin-top:24px;padding:10px;border:1px dashed #C9A86C;background:#FBF7EF;font-size:11px;color:#7A5C36}}
.assin{{margin-top:48px;text-align:center}}
@media print{{body{{margin:0}}}}
</style></head><body>{corpo}
<div class="aviso">⚠ MINUTA gerada automaticamente pelo EJC a partir da análise do extrato. Não constitui
afirmação de ilegalidade nem promessa de resultado. Revisão e validação obrigatórias do advogado responsável (OAB).</div>
</body></html>"""


def gerar_documento(tipo: str, analise: dict, cobrancas: list[dict], dados: dict | None = None) -> str:
    dados = dados or {}
    adv = html.escape(dados.get("advogado_nome") or "[ADVOGADO]")
    oab = html.escape(dados.get("advogado_oab") or "[OAB/UF nº ____]")
    cliente = html.escape(dados.get("cliente_nome") or analise.get("cliente_nome") or "[CLIENTE]")
    banco = html.escape(analise.get("banco") or "[INSTITUIÇÃO FINANCEIRA]")
    cidade = html.escape(dados.get("cidade") or "Betim/MG")
    total = float(analise.get("total_abusivo") or 0)
    hoje = date.today().strftime("%d/%m/%Y")
    tabela = (f"<table><tr><th>Data</th><th>Cobrança</th><th>Base legal</th><th>Valor</th></tr>"
              f"{_linhas_cobrancas_html(cobrancas)}</table>")

    if tipo == "notificacao":
        corpo = f"""<h1>Notificação Extrajudicial</h1>
<p><b>Notificante:</b> {cliente}.</p>
<p><b>Notificado:</b> {banco}.</p>
<p>Pela presente, o(a) notificante, por seu advogado, vem expor e ao final requerer:</p>
<p>Da análise do extrato bancário foram identificadas as seguintes cobranças <b>potencialmente indevidas</b>,
no valor total de <b>R$ {total:.2f}</b>, que se passa a discriminar:</p>
{tabela}
<p>Diante do exposto, <b>NOTIFICA-SE</b> a instituição para que, no prazo de <b>10 (dez) dias</b>, promova a
restituição/estorno dos valores apontados ou apresente justificativa documental, sob pena de adoção das
medidas judiciais cabíveis.</p>
<p>{cidade}, {hoje}.</p>
<div class="assin">_______________________________<br>{adv} — {oab}</div>"""
    elif tipo == "peticao":
        corpo = f"""<h1>Petição Inicial — Repetição de Indébito c/c Danos</h1>
<p>EXCELENTÍSSIMO(A) SENHOR(A) DOUTOR(A) JUIZ(A) DE DIREITO DA ___ VARA CÍVEL DA COMARCA DE {html.escape(dados.get('comarca') or '________')}.</p>
<p><b>{cliente}</b>, por seu advogado (procuração anexa), vem propor a presente <b>AÇÃO DE REPETIÇÃO DE INDÉBITO
c/c REVISÃO DE TARIFAS</b> em face de <b>{banco}</b>, pelos fatos e fundamentos a seguir.</p>
<p><b>DOS FATOS.</b> A análise do extrato revelou cobranças <b>potencialmente indevidas</b> totalizando
<b>R$ {total:.2f}</b>:</p>
{tabela}
<p><b>DO DIREITO.</b> As cobranças apontadas encontram questionamento na legislação consumerista e nas normas do
CMN/BCB indicadas no quadro acima — a serem confirmadas e fundamentadas pelo advogado.</p>
<p><b>DOS PEDIDOS.</b> Requer-se: (a) a declaração de inexigibilidade das cobranças indevidas; (b) a restituição
na forma do art. 42, parágrafo único, do CDC; (c) o que mais for de direito.</p>
<p>Dá-se à causa o valor de R$ {total:.2f}.</p>
<p>{cidade}, {hoje}.</p>
<div class="assin">_______________________________<br>{adv} — {oab}</div>"""
    elif tipo == "bacen":
        corpo = f"""<h1>Reclamação — Banco Central (RDR)</h1>
<p><b>Reclamante:</b> {cliente}.</p>
<p><b>Instituição reclamada:</b> {banco}.</p>
<p><b>Resumo:</b> identificação de cobranças possivelmente em desacordo com a regulamentação do CMN/BCB,
totalizando R$ {total:.2f}, conforme análise do extrato:</p>
{tabela}
<p><b>Pedido:</b> apuração pela instituição e pelo Banco Central das cobranças apontadas, com estorno dos
valores indevidos, nos termos da regulamentação aplicável.</p>
<p>{hoje}.</p>"""
    else:
        raise ValueError("tipo deve ser notificacao|peticao|bacen")
    return _doc_base(tipo, corpo)
