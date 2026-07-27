"""Regras de negócio da Sala Jurídica Conversacional (V1).

Fluxo central: mensagem do advogado → SingleAICoreOrchestrator (sanitização
LGPD, RAG, AILog, HITL — nada é contornado aqui) → persistência da resposta
como LegalChatMessage → nova versão do estado jurídico consolidado.
"""
from __future__ import annotations

import shutil
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any
from uuid import uuid4

from fastapi import HTTPException
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.client_ownership import (
    obter_cliente_autorizado,
    pode_ver_caso_resumido,
    pode_ver_cliente,
)
from app.core.config import get_settings
from app.core.security import ROLE_LEVEL
from app.models.case import Case, CaseArea, CaseStatus
from app.models.case_intelligence import CaseIntelligenceSnapshot
from app.models.client import Client
from app.models.document import DocConfidencialidade, Document
from app.models.legal_chat import (
    LegalChatAttachment,
    LegalChatMessage,
    LegalChatSession,
    LegalChatStateVersion,
)
from app.models.user import User
from app.schemas.legal_chat import ConverterRequest, MensagemCreate
from app.services.case_numeracao import proximo_numero_interno
# Núcleo de matching COMPARTILHADO de conflito (EOAB arts. 34-35) — mesma
# implementação usada por detectar_conflito/verificar_conflito; a Sala só
# FORMATA os achados no schema carteira-protegido (padrão Raio-X), sem criar
# uma terceira regra de correspondência.
from app.services.conflito_service import (
    _casos_por_parte_contraria,
    _clientes_por_nome,
)

# Modo do seletor → task_type do gateway. "chat_rapido" NÃO consta no mapa do
# intent_classifier de propósito: a conversa livre cai no fallback por
# keywords (agente ESPECIALISTA por área) ou no CaseAgent — ambos com
# tarefa_padrao séria (→ gateway "estrategia", tier alto). Não substituir por
# um task_type mapeado: isso pularia a seleção de especialista sem ganhar tier.
MODO_TASK_TYPE: dict[str, str] = {
    "conversa_livre": "chat_rapido",
    "organizar_fatos": "analise_juridica",
    "analisar_provas": "analise_juridica",
    "detectar_contradicoes": "analise_juridica",
    "estrategia_da_parte": "estrategia",
    "simular_defesa": "estrategia",
    "julgar_caso": "analise_juridica",
    "pesquisar_direito": "pesquisa_juridica",
    "elaborar_documento": "elaboracao_peca",
    "revisar_documento": "analise_juridica",
}

# Instrução de método por modo — complementa (não substitui) os prompts de
# ramo resolvidos pelo núcleo de IA.
MODO_INSTRUCAO: dict[str, str] = {
    # Conversa livre TAMBÉM tem método: profundidade de análise, estrutura e
    # honestidade epistêmica — o advogado espera o nível de um chat jurídico
    # de fronteira, não uma resposta telegráfica.
    "conversa_livre": (
        "Responda com profundidade técnica e didática: estruture em seções "
        "curtas quando a resposta passar de um parágrafo; fundamente cada "
        "afirmação jurídica (lei, súmula ou precedente VERIFICÁVEL — nunca "
        "invente); distinga expressamente o que é certo, o que é provável e o "
        "que depende de fato/documento ausente; aponte riscos e caminhos "
        "alternativos; termine com os próximos passos práticos quando fizer "
        "sentido. Se a pergunta for simples, seja direto — sem encher."
    ),
    "organizar_fatos": (
        "Organize os fatos em cronologia, distinguindo expressamente: "
        "comprovado, alegado, inferido, controvertido, ausente e superado."
    ),
    "analisar_provas": (
        "Relacione cada fato às provas existentes (documento, trecho, origem) "
        "e aponte a força probatória aparente e o que exige perícia."
    ),
    "detectar_contradicoes": (
        "Compare documentos, valores, datas, nomes, locais e versões; liste "
        "cada divergência com as fontes confrontadas."
    ),
    "estrategia_da_parte": (
        "Construa a melhor tese para o cliente: tese favorável, tese contrária "
        "provável, fragilidades, documentos faltantes e próximos passos."
    ),
    "simular_defesa": (
        "Atue como advogado da parte contrária: ataque os pontos frágeis e "
        "liste as impugnações prováveis, com fundamento."
    ),
    "julgar_caso": (
        "Analise como magistrado: pontos controvertidos, provas necessárias e "
        "provável distribuição do ônus probatório."
    ),
    "pesquisar_direito": (
        "Pesquise legislação e fontes oficiais; cite somente fontes "
        "verificáveis e sinalize expressamente o que ainda exige verificação."
    ),
    "elaborar_documento": (
        "Antes de redigir, confirme polo, objetivo, fase e prazo se não "
        "estiverem evidentes. Não insira dados fictícios: lacunas viram "
        "campos [PREENCHER]."
    ),
    "revisar_documento": (
        "Revise coerência, fatos, pedidos, fundamentação, competência, "
        "legitimidade, valores, dispositivos citados e contradições."
    ),
}

_GESTAO = {"superadmin", "admin", "socio"}
_CONVERSION_ROLES = _GESTAO | {"advogado"}


def _role(user: User) -> str:
    return getattr(user.role, "value", str(user.role))


async def obter_sessao(
    db: AsyncSession, session_id: str, user: User
) -> LegalChatSession:
    """Ownership fail-closed: criador, responsável ou gestão. 404 se inexistente."""
    sessao = await db.get(LegalChatSession, session_id)
    if sessao is None or sessao.deleted_at is not None:
        raise HTTPException(404, "Análise não encontrada")
    if _role(user) in _GESTAO:
        return sessao
    if user.id not in {sessao.created_by, sessao.advogado_responsavel_id}:
        raise HTTPException(403, "Sem acesso a esta análise")
    return sessao


def exigir_nao_congelada(sessao: LegalChatSession) -> None:
    if sessao.frozen_at is not None:
        raise HTTPException(
            409, "Análise convertida em caso está congelada para auditoria"
        )


async def ultima_versao_estado(
    db: AsyncSession, session_id: str
) -> LegalChatStateVersion | None:
    res = await db.execute(
        select(LegalChatStateVersion)
        .where(LegalChatStateVersion.session_id == session_id)
        .order_by(LegalChatStateVersion.versao.desc())
        .limit(1)
    )
    return res.scalar_one_or_none()


async def gravar_versao_estado(
    db: AsyncSession,
    sessao: LegalChatSession,
    *,
    estado: dict,
    resumo: str | None,
    origem: str,
    created_by: str | None,
) -> LegalChatStateVersion:
    # Lock da linha da sessão: serializa mensagens/edições CONCORRENTES da
    # mesma sessão. Sem ele, duas transações leem a mesma última versão e a
    # segunda viola uq_legal_chat_estado_versao (500 em vez de enfileirar).
    # populate_existing: força a releitura dos atributos sob o lock (o objeto
    # do identity map pode estar stale) — fecha o TOCTOU de congelamento com
    # uma conversão concorrente que congelou a sessão durante a chamada de IA.
    res = await db.execute(
        select(LegalChatSession)
        .where(LegalChatSession.id == sessao.id)
        .with_for_update()
        .execution_options(populate_existing=True)
    )
    travada = res.scalar_one_or_none()
    if travada is not None:
        exigir_nao_congelada(travada)
    atual = await ultima_versao_estado(db, sessao.id)
    versao = (atual.versao if atual else 0) + 1
    nova = LegalChatStateVersion(
        id=str(uuid4()),
        session_id=sessao.id,
        versao=versao,
        resumo=resumo,
        estado=estado,
        origem=origem,
        created_by=created_by,
    )
    db.add(nova)
    return nova


# Tetos dos blocos de contexto conversacional — protegem a janela de tokens
# do provider em sessões longas/com muitos anexos. Ampliados (2026-07): a Sala
# roteia para modelos de janela grande (200k tokens) e a memória curta era o
# principal fator de resposta "rasa" em sessões longas; ~48k chars ≈ 12k tokens
# continua folgado até para o provider local.
_HISTORICO_MAX_MENSAGENS = 20
_HISTORICO_MAX_CHARS_MSG = 3_000
_HISTORICO_MAX_CHARS_TOTAL = 24_000
_ANEXO_MAX_CHARS = 6_000
_ANEXOS_MAX_CHARS_TOTAL = 24_000


def _montar_mensagem_ia(
    payload: MensagemCreate,
    sessao: LegalChatSession,
    historico: list[LegalChatMessage] | None = None,
    anexos: list[LegalChatAttachment] | None = None,
) -> str:
    """Mensagem efetiva enviada ao núcleo: instrução do modo + workspace +
    histórico recente + síntese dos anexos + texto do advogado.

    Sem o histórico, cada mensagem chegava ao núcleo sem memória da conversa;
    sem os anexos, os documentos enviados nunca entravam no contexto da IA.
    """
    import json

    partes: list[str] = []
    instrucao = MODO_INSTRUCAO.get(payload.modo) or ""
    if instrucao:
        partes.append(f"[MÉTODO DO MODO '{payload.modo}'] {instrucao}")
    if payload.incluir_workspace and (sessao.workspace_texto or "").strip():
        partes.append(
            "[ÁREA DE TRABALHO DO ADVOGADO — fatos, anotações e rascunhos]\n"
            + sessao.workspace_texto.strip()
        )

    # Histórico (cronológico): truncagem por mensagem + teto total, cortando
    # as MAIS ANTIGAS primeiro (itera das novas para as velhas e reverte).
    linhas: list[str] = []
    total = 0
    for m in reversed(historico or []):
        autor = "Advogado" if m.autor == "user" else "IA"
        linha = f"{autor}: {(m.conteudo or '')[:_HISTORICO_MAX_CHARS_MSG]}"
        if total + len(linha) > _HISTORICO_MAX_CHARS_TOTAL:
            break
        linhas.append(linha)
        total += len(linha)
    if linhas:
        partes.append("[HISTÓRICO DA CONVERSA]\n" + "\n\n".join(reversed(linhas)))

    # Anexos: nome + síntese compacta da extração estruturada, com teto por
    # anexo e teto total.
    blocos_anexos: list[str] = []
    total = 0
    for a in anexos or []:
        sintese = json.dumps(
            a.resultado_analise or {}, ensure_ascii=False, separators=(",", ":")
        )[:_ANEXO_MAX_CHARS]
        bloco = f"- {a.nome_original}: {sintese}"
        if total + len(bloco) > _ANEXOS_MAX_CHARS_TOTAL:
            break
        blocos_anexos.append(bloco)
        total += len(bloco)
    if blocos_anexos:
        partes.append("[DOCUMENTOS ANEXADOS]\n" + "\n".join(blocos_anexos))

    partes.append(payload.conteudo)
    return "\n\n".join(partes)


async def enviar_mensagem(
    db: AsyncSession,
    sessao: LegalChatSession,
    payload: MensagemCreate,
    user: User,
) -> dict[str, Any]:
    """Persiste a pergunta, roda o núcleo único de IA e persiste a resposta."""
    exigir_nao_congelada(sessao)

    # Contexto conversacional carregado ANTES de persistir a nova pergunta —
    # assim a própria mensagem não entra duplicada no histórico enviado à IA.
    res_hist = await db.execute(
        select(LegalChatMessage)
        .where(LegalChatMessage.session_id == sessao.id)
        .order_by(LegalChatMessage.created_at.desc())
        .limit(_HISTORICO_MAX_MENSAGENS)
    )
    historico = list(reversed(res_hist.scalars().all()))  # ordem cronológica
    res_anexos = await db.execute(
        select(LegalChatAttachment)
        .where(LegalChatAttachment.session_id == sessao.id)
        .order_by(LegalChatAttachment.created_at)
    )
    anexos = list(res_anexos.scalars().all())

    msg_user = LegalChatMessage(
        id=str(uuid4()),
        session_id=sessao.id,
        autor="user",
        user_id=user.id,
        modo=payload.modo,
        conteudo=payload.conteudo,
    )
    db.add(msg_user)
    await db.flush()

    # Import tardio: mantém o service importável em testes sem stack de IA.
    from app.services.ai.core.orchestrator import run_ai_task

    resultado = await run_ai_task(
        db=db,
        user=user,
        task_type=MODO_TASK_TYPE[payload.modo],
        mensagem=_montar_mensagem_ia(payload, sessao, historico, anexos),
        case_id=sessao.convertido_case_id,
        params={
            "module_key": "sala-juridica",
            "surface": "sala_juridica",
            # Anexa o padrão obrigatório da Sala ao system prompt do agente
            # (system_prompts/sala_juridica.py) — nunca via mensagem do usuário.
            "prompt_extra": "sala_juridica",
        },
        usar_rag=payload.usar_rag,
    )

    custo = Decimal(str(resultado.get("custo_estimado_brl") or 0))
    msg_ia = LegalChatMessage(
        id=str(uuid4()),
        session_id=sessao.id,
        autor="ia",
        modo=payload.modo,
        conteudo=resultado.get("conteudo") or "",
        modelo=resultado.get("modelo"),
        agente=resultado.get("agente"),
        skills=list(resultado.get("skills_nativas") or []),
        fontes=list(resultado.get("fontes") or []),
        citacoes=list(resultado.get("citacoes") or []),
        alertas=list(resultado.get("alertas") or []),
        tokens_input=resultado.get("tokens_input"),
        tokens_output=resultado.get("tokens_output"),
        custo_estimado=custo,
        ai_log_id=resultado.get("log_id"),
    )
    db.add(msg_ia)

    # Estado: merge automático de fontes + (V2) extração estruturada pela IA
    # local. A extração NUNCA bloqueia a resposta: qualquer falha degrada para
    # o merge de fontes, e a curadoria fina segue disponível via PATCH /estado.
    atual = await ultima_versao_estado(db, sessao.id)
    estado = dict((atual.estado if atual else {}) or {})
    fontes = {(f.get("titulo"), f.get("fonte")): f for f in estado.get("fontes", [])}
    for f in resultado.get("fontes") or []:
        fontes.setdefault((f.get("titulo"), f.get("fonte")), f)
    estado["fontes"] = list(fontes.values())

    origem_estado = "ia"
    resumo_estado: str | None = None
    custo_extracao = Decimal("0")
    if get_settings().SALA_JURIDICA_AUTO_ESTADO:
        extraido, custo_extracao = await _extrair_estado_automatico(
            db, user, estado_atual=estado,
            pergunta=payload.conteudo,
            resposta=resultado.get("conteudo") or "",
        )
        if extraido is not None:
            resumo_estado = extraido.pop("_resumo", None)
            extraido["fontes"] = estado["fontes"]  # fontes vêm do RAG, não do LLM
            # Merge PARCIAL: o extrator pode devolver só algumas chaves (ex.:
            # apenas "fatos"). As omitidas herdam do estado atual — substituir
            # o dicionário inteiro apagaria provas/riscos/cronologia já
            # consolidados em versões anteriores.
            estado = {**estado, **extraido}
            origem_estado = "ia_extracao"

    versao = await gravar_versao_estado(
        db, sessao,
        estado=estado,
        resumo=resumo_estado,
        origem=origem_estado,
        created_by=None,
    )
    msg_ia.estado_versao = versao.versao
    # Custo total da sessão inclui TAMBÉM a chamada de extração de estado —
    # ela pode cair em fallback externo pago e não pode sumir da contabilidade.
    sessao.custo_ia_total = (
        (sessao.custo_ia_total or Decimal("0")) + custo + custo_extracao
    )
    await db.flush()

    return {
        "mensagem_user": serializar_mensagem(msg_user),
        "mensagem_ia": serializar_mensagem(msg_ia),
        "estado_versao": versao.versao,
        "is_rascunho": resultado.get("is_rascunho", True),
        "aviso_hitl": resultado.get("aviso_hitl"),
        "critica_adversarial": resultado.get("critica_adversarial"),
    }


_CHAVES_ESTADO = {
    "fatos", "provas", "contradicoes", "questoes", "teses",
    "riscos", "pendencias", "cronologia", "fontes",
}

_PROMPT_EXTRACAO = """Você é o extrator de estado jurídico da Sala Jurídica.
Atualize o ESTADO CONSOLIDADO abaixo com base na última interação, e responda
SOMENTE com um objeto JSON válido (sem markdown, sem comentários) com as chaves:
fatos, provas, contradicoes, questoes, teses, riscos, pendencias, cronologia,
_resumo (string de até 3 frases com a síntese atual).

Regras invioláveis:
- cada fato tem {{"texto": ..., "classificacao": "comprovado"|"alegado"|"inferido"|"controvertido"|"ausente"|"superado"}};
- NUNCA promova um fato a "comprovado" sem prova documental mencionada;
- fato substituído por informação posterior vira "superado" (não é apagado);
- riscos têm {{"descricao": ..., "nivel": "baixo"|"medio"|"alto"}};
- não invente fatos, provas nem fontes que não constem da interação/estado.

ESTADO CONSOLIDADO ATUAL:
{estado}

PERGUNTA DO ADVOGADO:
{pergunta}

RESPOSTA DA ANÁLISE:
{resposta}"""


async def _extrair_estado_automatico(
    db,
    user,
    *,
    estado_atual: dict,
    pergunta: str,
    resposta: str,
) -> tuple[dict | None, Decimal]:
    """Extração estruturada do estado via task_type "resumo".

    Retorna (novo estado validado | None, custo estimado em BRL). O estado é
    None em qualquer falha (fail-soft): JSON inválido, chaves desconhecidas,
    provider indisponível. O custo já incorrido é SEMPRE devolvido: a cadeia
    do task_type "resumo" prioriza o provider local (ollama), mas pode cair
    em fallback externo pago (maritaca/groq) e o orchestrator não expõe uma
    forma de forçar rota local-only por chamada — então o gasto é
    contabilizado em custo_ia_total pelo chamador.
    """
    import json

    from app.services.ai.core.orchestrator import run_ai_task

    custo = Decimal("0")
    try:
        mensagem = _PROMPT_EXTRACAO.format(
            estado=json.dumps(
                {k: v for k, v in estado_atual.items() if k != "fontes"},
                ensure_ascii=False,
            )[:12_000],
            pergunta=pergunta[:4_000],
            resposta=resposta[:12_000],
        )
        r = await run_ai_task(
            db=db,
            user=user,
            task_type="resumo",
            mensagem=mensagem,
            params={"module_key": "sala-juridica", "surface": "sala_juridica_estado"},
            usar_rag=False,
        )
        custo = Decimal(str(r.get("custo_estimado_brl") or 0))
        bruto = r.get("conteudo") or ""
        inicio, fim = bruto.find("{"), bruto.rfind("}")
        if inicio < 0 or fim <= inicio:
            return None, custo
        dados = json.loads(bruto[inicio : fim + 1])
        if not isinstance(dados, dict):
            return None, custo
        resumo = dados.pop("_resumo", None)
        if set(dados) - _CHAVES_ESTADO:
            return None, custo
        if not all(isinstance(v, list) for v in dados.values()):
            return None, custo
        if resumo is not None:
            dados["_resumo"] = str(resumo)[:2_000]
        return dados, custo
    except Exception:  # fail-soft deliberado: extração jamais bloqueia a resposta
        return None, custo


# ── Preview de conversão (paridade com o Raio-X) ────────────────────────────
# A Sala é a porta de entrada principal do escritório: a conversão em caso não
# pode depender só da autodeclaração do advogado. O preview cruza a base por
# dever ético (EOAB arts. 34-35) SEM expor outra carteira (padrão Raio-X:
# correspondência em carteira não autorizada vira alerta "protegido").

def _alerta_protegido(tipo: str, mensagem: str) -> dict[str, Any]:
    return {
        "tipo": tipo,
        "nome": "Correspondência protegida na base do escritório",
        "mensagem": mensagem,
        "protegido": True,
        "confirmado": False,
    }


def _nomes_partes_anexos(anexos: list[LegalChatAttachment]) -> list[str]:
    """Nomes de partes extraídos pela análise dos anexos (fail-soft)."""
    nomes: list[str] = []
    vistos: set[str] = set()
    for a in anexos:
        ra = a.resultado_analise if isinstance(a.resultado_analise, dict) else {}
        intake = ra.get("intake_result") if isinstance(ra.get("intake_result"), dict) else ra
        for item in intake.get("partes") or []:
            if isinstance(item, str):
                nome = item.strip()
            elif isinstance(item, dict):
                bruto = item.get("nome") or item.get("valor") or item.get("parte") or ""
                if isinstance(bruto, dict):
                    bruto = bruto.get("valor") or ""
                nome = str(bruto).strip()
            else:
                continue
            chave = nome.lower()
            # Piso de 4 chars (mesmo das primitivas de matching) evita ILIKE espúrio.
            if len(nome) >= 4 and chave not in vistos:
                vistos.add(chave)
                nomes.append(nome)
    return nomes[:20]


async def preview_conversao(
    db: AsyncSession,
    sessao: LegalChatSession,
    user: User,
    *,
    nome_cliente: str | None = None,
    client_id: str | None = None,
) -> dict[str, Any]:
    """Conflitos + duplicados ANTES da conversão, sem vazar carteira alheia."""
    res = await db.execute(
        select(LegalChatAttachment)
        .where(LegalChatAttachment.session_id == sessao.id)
        .order_by(LegalChatAttachment.created_at)
    )
    anexos = list(res.scalars().all())
    partes = _nomes_partes_anexos(anexos)
    candidato = (nome_cliente or sessao.cliente_potencial or "").strip()
    # Cliente EXISTENTE selecionado: o nome canônico dele é o candidato — sem
    # isso, sessões sem cliente_potencial pulavam a checagem de o cliente ser
    # parte contrária em outro caso do escritório (apontamento Codex P1).
    # Sob o gate canônico: UUID de outra carteira devolve 404 aqui, em vez de
    # virar sonda de identidade/contagem de casos alheios.
    if client_id:
        cli = await obter_cliente_autorizado(db, user, client_id)
        if not candidato:
            candidato = (cli.nome_exibicao or "").strip()
    candidato_lower = candidato.lower()

    alertas: list[dict[str, Any]] = []
    # 1) Parte dos anexos já é cliente do escritório (atual ou anterior).
    for nome in partes:
        if nome.lower() == candidato_lower:
            continue
        for c in await _clientes_por_nome(db, nome, limit=5):
            if await pode_ver_cliente(db, user, c):
                alertas.append({
                    "tipo": "parte_corresponde_a_cliente",
                    "nome": c.nome_exibicao,
                    "client_id": c.id,
                    "mensagem": "Parte dos documentos possui nome semelhante a cliente atual ou anterior. Revisar conflito de interesses.",
                    "protegido": False,
                    "confirmado": False,
                })
            else:
                alertas.append(_alerta_protegido(
                    "parte_corresponde_a_cliente_protegido",
                    "Correspondência em carteira protegida. Solicite revisão de conflito à gestão antes de prosseguir.",
                ))
    # 2) Parte dos anexos figura como parte contrária em caso do escritório;
    # 3) o PRÓPRIO cliente candidato é parte contrária em caso nosso (grave).
    for nome, grave in [(n, False) for n in partes] + ([(candidato, True)] if len(candidato) >= 4 else []):
        for caso in await _casos_por_parte_contraria(db, nome, limit=5):
            if pode_ver_caso_resumido(user, caso):
                alertas.append({
                    "tipo": ("CONFLITO_cliente_e_parte_contraria" if grave
                             else "parte_corresponde_a_parte_contraria"),
                    "nome": caso.parte_contraria,
                    "case_id": caso.id,
                    "mensagem": ("O cliente informado figura como PARTE CONTRÁRIA em caso do escritório — conflito potencial grave (EOAB art. 17/34)."
                                 if grave else
                                 "Parte dos documentos possui nome semelhante a parte contrária de caso acessível. Revisar conflito e grupo econômico."),
                    "protegido": False,
                    "confirmado": False,
                })
            else:
                alertas.append(_alerta_protegido(
                    "parte_corresponde_a_caso_protegido",
                    "Correspondência em caso protegido. Solicite revisão de conflito à gestão antes de prosseguir.",
                ))
    # Dedup sem depender de ids protegidos (mesma chave do Raio-X).
    unicos: list[dict[str, Any]] = []
    vistos: set[tuple[str, str, str]] = set()
    for a in alertas:
        k = (str(a.get("tipo") or ""), str(a.get("nome") or ""), str(a.get("mensagem") or ""))
        if k not in vistos:
            vistos.add(k)
            unicos.append(a)

    # Clientes possivelmente duplicados (só relevante ao CRIAR cliente novo).
    # Os protegidos colapsam numa ÚNICA entrada: a contagem de correspondências
    # em carteira alheia também é informação (mesma doutrina de alertas_conflito).
    clientes_dup: list[dict[str, Any]] = []
    if candidato and not client_id:
        houve_protegido = False
        for c in await _clientes_por_nome(db, candidato, limit=10):
            if await pode_ver_cliente(db, user, c):
                clientes_dup.append({
                    "id": c.id, "nome": c.nome_exibicao, "protegido": False,
                })
            else:
                houve_protegido = True
        if houve_protegido:
            clientes_dup.append({
                "id": None,
                "nome": "Cliente protegido na base do escritório",
                "protegido": True,
            })

    # Casos ATIVOS do cliente escolhido — sinal de caso possivelmente duplicado.
    # Só roda com client_id JÁ autorizado (obter_cliente_autorizado acima);
    # protegidos colapsam numa entrada para não vazar a contagem.
    casos_ativos: list[dict[str, Any]] = []
    if client_id:
        rows = (await db.execute(
            select(Case).where(
                Case.client_id == client_id,
                Case.deleted_at.is_(None),
                Case.status.notin_([CaseStatus.encerrado, CaseStatus.arquivado]),
            ).limit(10)
        )).scalars().all()
        houve_protegido = False
        for caso in rows:
            if pode_ver_caso_resumido(user, caso):
                casos_ativos.append({
                    "id": caso.id,
                    "titulo": caso.titulo,
                    "numero_interno": caso.numero_interno,
                    "protegido": False,
                })
            else:
                houve_protegido = True
        if houve_protegido:
            casos_ativos.append({
                "id": None, "titulo": "Caso protegido", "numero_interno": None,
                "protegido": True,
            })

    return {
        "session_id": sessao.id,
        "alertas_conflito": unicos,
        "clientes_possivelmente_duplicados": clientes_dup,
        "casos_ativos_do_cliente": casos_ativos,
        "anexos_disponiveis": [
            {"id": a.id, "nome": a.nome_original, "tipo": a.tipo_documento}
            for a in anexos
        ],
        "bloqueia": bool(unicos),
    }


def _copiar_anexo(
    anexo: LegalChatAttachment, case_id: str, client_id: str, user: User,
) -> tuple[Document, Path]:
    """Copia o arquivo físico do anexo e devolve (Document oficial, caminho
    copiado — para limpeza em caso de rollback)."""
    settings = get_settings()
    origem = Path(settings.UPLOAD_DIR) / anexo.filepath
    rel = Path("sala-juridica-convertidos") / case_id / f"{uuid4()}{origem.suffix}"
    destino = Path(settings.UPLOAD_DIR) / rel
    destino.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(origem, destino)
    # Texto extraído/sanitizado preservado no upload → Document pesquisável no
    # GED e disponível ao context builder (apontamento Codex P2). Anexos
    # antigos (antes da mudança no upload) seguem sem texto — fail-soft.
    ra = anexo.resultado_analise if isinstance(anexo.resultado_analise, dict) else {}
    texto = ra.get("_texto_sanitizado")
    doc = Document(
        id=str(uuid4()),
        titulo=anexo.nome_original[:255],
        descricao="Documento transferido da análise da Sala Jurídica.",
        tipo=anexo.tipo_documento or "outro",
        filename=anexo.nome_original,
        filepath=str(rel),
        mimetype=anexo.mimetype,
        size_bytes=anexo.size_bytes,
        ocr_text=(str(texto)[:200_000]
                  if isinstance(texto, str) and texto.strip() else None),
        confidencialidade=DocConfidencialidade.interno,
        case_id=case_id,
        client_id=client_id,
        uploaded_by=user.id,
    )
    return doc, destino


async def _transferir_anexos(
    db: AsyncSession, sessao: LegalChatSession, case_id: str, client_id: str, user: User,
) -> tuple[list[str], list[Path]]:
    """Todos os anexos da sessão → Document do caso.

    FAIL-CLOSED (apontamento Codex P1): arquivo físico ausente ABORTA a
    conversão ANTES do congelamento — a sessão fica retryável e o caso nunca
    nasce sem documento fundante. Retorna também os caminhos copiados, para o
    chamador limpar em caso de falha posterior na transação."""
    res = await db.execute(
        select(LegalChatAttachment)
        .where(LegalChatAttachment.session_id == sessao.id)
        .order_by(LegalChatAttachment.created_at)
    )
    anexos = list(res.scalars().all())
    settings = get_settings()
    faltantes = [
        a.nome_original for a in anexos
        if not (Path(settings.UPLOAD_DIR) / a.filepath).is_file()
    ]
    if faltantes:
        raise HTTPException(422, {
            "mensagem": ("Arquivo físico indisponível para transferência. "
                         "Restaure o storage e tente novamente, ou converta "
                         "sem transferir anexos."),
            "anexos_indisponiveis": faltantes,
        })
    transferidos: list[str] = []
    copiados: list[Path] = []
    for anexo in anexos:
        doc, caminho = _copiar_anexo(anexo, case_id, client_id, user)
        db.add(doc)
        transferidos.append(doc.id)
        copiados.append(caminho)
    return transferidos, copiados


def _limpar_copias(copiados: list[Path]) -> None:
    """Best-effort: remove cópias físicas órfãs quando a transação falha."""
    for caminho in copiados:
        try:
            caminho.unlink(missing_ok=True)
        except OSError:
            pass


async def _criar_snapshot_sala(
    db: AsyncSession, sessao: LegalChatSession, case_id: str, area: str, user: User,
) -> int:
    """CaseIntelligenceSnapshot origem 'sala_juridica' com o estado consolidado.
    Versão = max(versao do caso)+1, serializado por advisory lock transacional
    por caso — dois vínculos concorrentes de sessões DIFERENTES ao mesmo caso
    não leem o mesmo max (o índice único uq_cis_case_versao segue como última
    linha de defesa). Retorna a versão criada."""
    estado_v = await ultima_versao_estado(db, sessao.id)
    estado = dict((estado_v.estado if estado_v else {}) or {})
    # Rastreabilidade: os logs mais RECENTES formaram o estado consolidado —
    # em sessões longas, cortar os antigos, nunca os novos (Codex P2).
    res_logs = await db.execute(
        select(LegalChatMessage.ai_log_id)
        .where(
            LegalChatMessage.session_id == sessao.id,
            LegalChatMessage.ai_log_id.isnot(None),
        )
        .order_by(LegalChatMessage.created_at.desc())
        .limit(100)
    )
    ai_log_ids = list(reversed(
        [v for v in res_logs.scalars().all() if isinstance(v, str)]
    ))
    await db.execute(
        text("SELECT pg_advisory_xact_lock(hashtext(:chave))"),
        {"chave": f"cis:{case_id}"},
    )
    proxima = (
        await db.scalar(
            select(func.coalesce(func.max(CaseIntelligenceSnapshot.versao), 0))
            .where(CaseIntelligenceSnapshot.case_id == case_id)
        )
    ) or 0
    snap = CaseIntelligenceSnapshot(
        id=str(uuid4()),
        case_id=case_id,
        versao=proxima + 1,
        origem="sala_juridica",
        payload={
            "area": area,
            "fatos": estado.get("fatos") or [],
            "provas": estado.get("provas") or [],
            "contradicoes": estado.get("contradicoes") or [],
            "questoes": estado.get("questoes") or [],
            "teses": {"principal": None, "secundarias": estado.get("teses") or []},
            "riscos": estado.get("riscos") or [],
            "pendencias": estado.get("pendencias") or [],
            "cronologia": estado.get("cronologia") or [],
            "fontes": estado.get("fontes") or [],
            "sala_juridica_session_id": sessao.id,
            "revisao_humana_obrigatoria": True,
        },
        resumo=((estado_v.resumo if estado_v else None) or None),
        ai_log_ids=ai_log_ids,
        criado_por=user.id,
        congelado=False,
        aprovado_por=None,
        aprovado_em=None,
    )
    db.add(snap)
    return snap.versao


async def converter_em_caso(
    db: AsyncSession,
    sessao: LegalChatSession,
    payload: ConverterRequest,
    user: User,
) -> dict[str, Any]:
    """Conversão controlada: cliente → conflito confirmado → caso → congelamento."""
    if _role(user) not in _CONVERSION_ROLES:
        raise HTTPException(
            403, "Seu perfil não possui autorização para criar casos oficiais"
        )
    # Lock pessimista da sessão + recheck: duas conversões simultâneas da
    # mesma análise não podem criar dois casos/clientes — a segunda transação
    # espera o lock e cai na idempotência abaixo. populate_existing garante
    # que o recheck lê os atributos recomprometidos, não o identity map stale.
    res = await db.execute(
        select(LegalChatSession)
        .where(LegalChatSession.id == sessao.id)
        .with_for_update()
        .execution_options(populate_existing=True)
    )
    sessao = res.scalar_one()
    if sessao.convertido_case_id:
        return {"case_id": sessao.convertido_case_id, "ja_convertido": True}
    if payload.area not in {a.value for a in CaseArea}:
        raise HTTPException(422, "Área do caso inválida")
    if bool(payload.client_id) == bool(payload.novo_cliente_nome):
        raise HTTPException(
            422, "Informe client_id OU novo_cliente_nome (exatamente um)"
        )

    # Gates de servidor (paridade Raio-X): conflito/duplicado DETECTADO exige
    # reconhecimento explícito do achado — o checkbox genérico não basta.
    preview = await preview_conversao(
        db, sessao, user,
        nome_cliente=payload.novo_cliente_nome,
        client_id=payload.client_id,
    )
    if preview["alertas_conflito"] and not payload.conflict_confirmed:
        raise HTTPException(409, {
            "mensagem": "Há alertas de conflito de interesses; revise os achados e confirme antes de converter.",
            "alertas_conflito": preview["alertas_conflito"],
        })
    if payload.novo_cliente_nome and preview["clientes_possivelmente_duplicados"] \
            and not payload.duplicate_confirmed:
        raise HTTPException(409, {
            "mensagem": "Há cliente possivelmente duplicado; selecione o existente ou confirme a criação.",
            "clientes_possivelmente_duplicados": preview["clientes_possivelmente_duplicados"],
        })
    # Cliente EXISTENTE com caso ativo: possível caso duplicado — o gate vale
    # para os dois caminhos, não só para cliente novo (apontamento Codex P1).
    if payload.client_id and preview["casos_ativos_do_cliente"] \
            and not payload.duplicate_confirmed:
        raise HTTPException(409, {
            "mensagem": "O cliente possui caso ativo; confirme que não se trata de caso duplicado antes de criar outro.",
            "casos_ativos_do_cliente": preview["casos_ativos_do_cliente"],
        })

    # Responsável: UUID livre no payload ia direto para o modelo — id
    # inexistente virava IntegrityError (500) e qualquer id permitia atribuir
    # caso/cliente a terceiro. Exige usuário ativo com piso de advogado.
    responsavel = await db.get(User, payload.advogado_responsavel_id)
    if (
        responsavel is None
        or not getattr(responsavel, "is_active", True)
        or ROLE_LEVEL.get(_role(responsavel), 0) < ROLE_LEVEL["advogado"]
    ):
        raise HTTPException(
            422, "Responsável inválido: informe um advogado ativo do escritório"
        )

    if payload.client_id:
        # Gate canônico (mesmo do Raio-X): 404 uniforme quando o UUID pertence
        # a outra carteira. Sem ele, converter com client_id alheio criava um
        # Case com o requisitante como responsável — e `pode_ver_cliente` passa
        # a conceder acesso permanente ao cliente por esse próprio vínculo
        # (escalonamento de carteira, não só escrita indevida).
        client = await obter_cliente_autorizado(db, user, payload.client_id)
    else:
        client = Client(
            id=str(uuid4()),
            nome=payload.novo_cliente_nome,
            responsavel_id=payload.advogado_responsavel_id,
        )
        db.add(client)
        await db.flush()

    numero = await proximo_numero_interno(db)
    case = Case(
        id=str(uuid4()),
        numero_interno=numero,
        titulo=payload.titulo_caso,
        area=CaseArea(payload.area),
        descricao_fatos=payload.descricao,
        client_id=client.id,
        advogado_responsavel_id=payload.advogado_responsavel_id,
        # G1 (mesma guarda de cases.py): caso em triagem nunca nasce sem
        # "o que fazer agora" — default aponta a revisão da análise convertida.
        proxima_acao=payload.proxima_acao
        or "Revisar a análise convertida da Sala Jurídica e definir a próxima providência",
    )
    db.add(case)
    await db.flush()

    # A inteligência preliminar acompanha o caso oficial (paridade Raio-X):
    # anexos viram Document e o estado consolidado vira snapshot versionado.
    # Falha em QUALQUER passo → cópias físicas removidas antes de propagar
    # (o banco faz rollback; arquivo órfão em disco não — Codex P2).
    transferidos: list[str] = []
    copiados: list[Path] = []
    try:
        if payload.transferir_anexos:
            transferidos, copiados = await _transferir_anexos(
                db, sessao, case.id, client.id, user
            )
        snapshot_versao = await _criar_snapshot_sala(
            db, sessao, case.id, payload.area, user
        )

        agora = datetime.now(timezone.utc)
        sessao.client_id = client.id
        sessao.convertido_case_id = case.id
        sessao.converted_at = agora
        sessao.frozen_at = agora  # congelada para auditoria — imutável daqui em diante
        sessao.status = "convertida_em_caso"
        await db.flush()
    except Exception:
        _limpar_copias(copiados)
        raise
    return {
        "case_id": case.id,
        "client_id": client.id,
        "ja_convertido": False,
        "documentos_transferidos": transferidos,
        "snapshot_versao": snapshot_versao,
    }


async def vincular_caso_existente(
    db: AsyncSession,
    sessao: LegalChatSession,
    case_id: str,
    user: User,
    *,
    transferir_anexos: bool = True,
) -> dict[str, Any]:
    """Vincula a análise a um caso JÁ EXISTENTE (sem criar caso novo).

    Mesmo desfecho de auditoria da conversão: sessão congelada e imutável.
    Ownership do caso é fail-closed (verificar_acesso_caso → 403/404).
    """
    if _role(user) not in _CONVERSION_ROLES:
        raise HTTPException(
            403, "Seu perfil não possui autorização para vincular a casos oficiais"
        )
    # Lock pessimista da sessão + recheck ANTES do gate de congelamento: a
    # repetição do vínculo já consumado é idempotente (ja_convertido=True em
    # vez de 409), e dois vínculos simultâneos não gravam case_ids distintos.
    # populate_existing garante recheck sobre atributos recomprometidos.
    res = await db.execute(
        select(LegalChatSession)
        .where(LegalChatSession.id == sessao.id)
        .with_for_update()
        .execution_options(populate_existing=True)
    )
    sessao = res.scalar_one()
    if sessao.convertido_case_id:
        return {"case_id": sessao.convertido_case_id, "ja_convertido": True}
    exigir_nao_congelada(sessao)

    from app.core.ownership import verificar_acesso_caso

    await verificar_acesso_caso(db, user, case_id)
    case = await db.get(Case, case_id)
    if case is None:
        raise HTTPException(404, "Caso não encontrado")

    # Mesmos efeitos da conversão: anexos → Document e estado → snapshot,
    # com a mesma limpeza de cópias físicas em falha.
    transferidos: list[str] = []
    copiados: list[Path] = []
    try:
        if transferir_anexos:
            transferidos, copiados = await _transferir_anexos(
                db, sessao, case.id, case.client_id, user
            )
        snapshot_versao = await _criar_snapshot_sala(
            db, sessao, case.id, getattr(case.area, "value", str(case.area)), user
        )

        agora = datetime.now(timezone.utc)
        sessao.client_id = case.client_id
        sessao.convertido_case_id = case.id
        sessao.converted_at = agora
        sessao.frozen_at = agora
        sessao.status = "convertida_em_caso"
        await db.flush()
    except Exception:
        _limpar_copias(copiados)
        raise
    return {
        "case_id": case.id,
        "client_id": case.client_id,
        "ja_convertido": False,
        "documentos_transferidos": transferidos,
        "snapshot_versao": snapshot_versao,
    }


# ── Exportação (DOCX/PDF) ────────────────────────────────────────────────────

_AVISO_EXPORT = (
    "Documento de trabalho gerado pela Sala Jurídica do EJC com apoio de IA. "
    "Conteúdo em rascunho, sujeito a revisão do advogado responsável (HITL)."
)


def _blocos_exportacao(
    sessao: LegalChatSession,
    mensagens: list[LegalChatMessage],
    estado: LegalChatStateVersion | None,
) -> list[tuple[str, str]]:
    """Blocos (título, corpo) comuns aos dois formatos de exportação."""
    blocos: list[tuple[str, str]] = []
    if (sessao.workspace_texto or "").strip():
        blocos.append(("Área de trabalho do advogado", sessao.workspace_texto.strip()))
    if estado is not None and estado.resumo:
        blocos.append((f"Síntese do estado jurídico (v{estado.versao})", estado.resumo))
    for m in mensagens:
        autor = "Advogado" if m.autor == "user" else f"Análise da IA ({m.modo})"
        blocos.append((autor, m.conteudo or ""))
    return blocos


def exportar_docx(
    sessao: LegalChatSession,
    mensagens: list[LegalChatMessage],
    estado: LegalChatStateVersion | None,
) -> bytes:
    import io

    from docx import Document

    doc = Document()
    doc.add_heading(sessao.titulo or "Sala Jurídica", level=0)
    doc.add_paragraph(_AVISO_EXPORT)
    for titulo, corpo in _blocos_exportacao(sessao, mensagens, estado):
        doc.add_heading(titulo, level=2)
        for par in corpo.split("\n\n"):
            doc.add_paragraph(par)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def exportar_pdf(
    sessao: LegalChatSession,
    mensagens: list[LegalChatMessage],
    estado: LegalChatStateVersion | None,
) -> bytes:
    from html import escape

    from weasyprint import HTML

    partes = [
        "<style>body{font-family:sans-serif;font-size:11pt;margin:2cm}"
        "h1{font-size:16pt}h2{font-size:12pt;margin-top:14pt;color:#7a5c14}"
        "p{white-space:pre-wrap;line-height:1.4}.aviso{font-size:8pt;color:#666}</style>",
        f"<h1>{escape(sessao.titulo or 'Sala Jurídica')}</h1>",
        f"<p class='aviso'>{escape(_AVISO_EXPORT)}</p>",
    ]
    for titulo, corpo in _blocos_exportacao(sessao, mensagens, estado):
        partes.append(f"<h2>{escape(titulo)}</h2><p>{escape(corpo)}</p>")
    return HTML(string="".join(partes)).write_pdf()


def serializar_mensagem(m: LegalChatMessage) -> dict[str, Any]:
    return {
        "id": m.id,
        "autor": m.autor,
        "modo": m.modo,
        "conteudo": m.conteudo,
        "modelo": m.modelo,
        "agente": m.agente,
        "skills": m.skills or [],
        "fontes": m.fontes or [],
        "citacoes": m.citacoes or [],
        "alertas": m.alertas or [],
        "tokens_input": m.tokens_input,
        "tokens_output": m.tokens_output,
        "custo_estimado": float(m.custo_estimado) if m.custo_estimado is not None else None,
        "ai_log_id": m.ai_log_id,
        "estado_versao": m.estado_versao,
        "created_at": m.created_at.isoformat() if m.created_at else None,
    }


def serializar_anexo(a: LegalChatAttachment) -> dict[str, Any]:
    return {
        "id": a.id,
        "nome_original": a.nome_original,
        "mimetype": a.mimetype,
        "size_bytes": a.size_bytes,
        "sha256": a.sha256,
        "tipo_documento": a.tipo_documento,
        "ocr_utilizado": a.ocr_utilizado,
        # Chaves `_`-prefixadas são internas por convenção do repo (o texto
        # sanitizado é retido só para virar Document.ocr_text na conversão) —
        # nunca saem na API, aqui como no caminho gêmeo de documentos.
        "resultado_analise": {
            k: v for k, v in (a.resultado_analise or {}).items()
            if not str(k).startswith("_")
        },
        "created_at": a.created_at.isoformat() if a.created_at else None,
    }


def serializar_sessao(
    s: LegalChatSession,
    *,
    incluir_relacionados: bool = False,
    estado: LegalChatStateVersion | None = None,
) -> dict[str, Any]:
    base: dict[str, Any] = {
        "id": s.id,
        "titulo": s.titulo,
        "status": s.status,
        "favorita": s.favorita,
        "cliente_potencial": s.cliente_potencial,
        "area_sugerida": s.area_sugerida,
        "advogado_responsavel_id": s.advogado_responsavel_id,
        "workspace_versao": s.workspace_versao,
        "client_id": s.client_id,
        "convertido_case_id": s.convertido_case_id,
        "frozen": s.frozen_at is not None,
        "custo_ia_total": float(s.custo_ia_total or 0),
        "created_by": s.created_by,
        "created_at": s.created_at.isoformat() if s.created_at else None,
        "updated_at": s.updated_at.isoformat() if s.updated_at else None,
    }
    if incluir_relacionados:
        base["workspace_texto"] = s.workspace_texto
        base["mensagens"] = [serializar_mensagem(m) for m in s.mensagens]
        base["anexos"] = [serializar_anexo(a) for a in s.anexos]
        base["estado"] = (
            {
                "versao": estado.versao,
                "resumo": estado.resumo,
                "estado": estado.estado,
                "origem": estado.origem,
                "created_at": estado.created_at.isoformat() if estado.created_at else None,
            }
            if estado
            else None
        )
    return base
