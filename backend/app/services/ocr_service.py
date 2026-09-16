# ── app/services/ocr_service.py ──────────────────────────────────────────────
# Extração de texto de documentos para busca por conteúdo.
# PDF nativo: PyMuPDF (rápido). PDF escaneado/imagens: Tesseract (se instalado).
# XML (NF-e): parse seguro (defusedxml se instalado; senão stdlib com bloqueio
# de DTD/entidades) + extração de campos fiscais estruturados.
# Falha de OCR NUNCA quebra o upload — apenas registra log.
from __future__ import annotations
import logging
import re

logger = logging.getLogger("ejc.ocr")

# Imports lazy/opcionais — sistema funciona sem OCR de imagem
try:
    import fitz  # PyMuPDF
    _PDF_OK = True
except ImportError:
    _PDF_OK = False

try:
    import pytesseract
    from PIL import Image
    _TESS_OK = True
except ImportError:
    _TESS_OK = False

try:
    import docx as _docx  # python-docx
    _DOCX_OK = True
except ImportError:
    _DOCX_OK = False

# XML: defusedxml (proteção XXE/billion-laughs nativa) com fallback stdlib.
# No fallback, XMLs com DTD/ENTITY são REJEITADOS antes do parse (ver _xml_root).
try:
    from defusedxml import ElementTree as _ET  # type: ignore
    _DEFUSED = True
except ImportError:
    import xml.etree.ElementTree as _ET  # noqa: N814
    _DEFUSED = False

MAX_OCR_CHARS = 200_000  # limite de armazenamento por documento


def extrair_texto(
    filepath: str,
    mimetype: str | None,
    suffix_hint: str | None = None,
) -> str | None:
    """
    Retorna texto extraído (truncado) ou None se tipo não suportado/falha.
    Síncrono — chamado via asyncio.to_thread no router.

    ``suffix_hint`` preserva a extensão já validada quando o arquivo físico
    ainda está em staging (ex.: ``.docx.uploading``). O hint só decide qual
    extrator local usar; nunca altera o path aberto.
    """
    mt = (mimetype or "").lower()
    hint = (suffix_hint or "").lower()
    try:
        if "pdf" in mt and _PDF_OK:
            return _pdf(filepath)
        if mt.startswith("image/") and _TESS_OK:
            return _imagem(filepath)
        if (
            "wordprocessingml" in mt
            or hint == ".docx"
            or filepath.lower().endswith(".docx")
        ) and _DOCX_OK:
            return _docx_texto(filepath)
        # Planilhas .xlsx via openpyxl (P1 2026-07-05: antes ocr_text ficava
        # NULL silencioso).
        if (
            "spreadsheetml" in mt
            or hint == ".xlsx"
            or filepath.lower().endswith(".xlsx")
        ):
            return _xlsx(filepath)
        if "ms-excel" in mt or filepath.lower().endswith(".xls"):
            return _xls_legado(filepath)
        if "msword" in mt or filepath.lower().endswith(".doc"):
            return _doc_legado(filepath)
        # XML antes do ramo text/ (text/xml também começa com "text/")
        if "xml" in mt or filepath.lower().endswith(".xml"):
            res = extrair_xml(filepath)
            return res.get("texto") if res else None
        # Markdown: texto plano (libmagic pode reportar text/plain ou
        # text/markdown) — leitura UTF-8 direta, como .txt (auditoria 12/08/2026).
        if mt.startswith("text/") or filepath.lower().endswith((".txt", ".md")):
            with open(filepath, encoding="utf-8", errors="ignore") as f:
                return f.read()[:MAX_OCR_CHARS]
    except Exception as e:
        logger.warning(f"OCR falhou para {filepath}: {e}")
    return None


def _pdf(path: str) -> str:
    doc = fitz.open(path)
    partes = []
    for page in doc:
        txt = page.get_text("text")
        # Página sem texto nativo (escaneada) → tentar Tesseract na imagem
        if len(txt.strip()) < 20 and _TESS_OK:
            pix = page.get_pixmap(dpi=200)
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            txt = pytesseract.image_to_string(img, lang="por")
        partes.append(txt)
        if sum(len(p) for p in partes) > MAX_OCR_CHARS:
            break
    doc.close()
    return "\n".join(partes)[:MAX_OCR_CHARS]


def _imagem(path: str) -> str:
    return pytesseract.image_to_string(Image.open(path), lang="por")[:MAX_OCR_CHARS]


def _docx_texto(path: str) -> str:
    # File-like evita qualquer dependência da extensão do path de staging.
    with open(path, "rb") as handle:
        d = _docx.Document(handle)
        return "\n".join(p.text for p in d.paragraphs)[:MAX_OCR_CHARS]


def _xlsx(path: str) -> str:
    """Texto pesquisável de planilha .xlsx: células não-vazias concatenadas
    por linha (' | ') e por aba ([Planilha: nome]), truncado em MAX_OCR_CHARS.
    read_only+data_only: streaming (não carrega o workbook inteiro) e valores
    calculados em vez de fórmulas. Erros sobem para o try de extrair_texto
    (log + None, padrão do arquivo — falha de extração nunca quebra upload)."""
    import openpyxl  # lazy, como os demais extratores

    # File-like evita a validação de extensão do openpyxl; o staging mantém
    # deliberadamente o sufixo ``.uploading`` até a promoção transacional.
    handle = open(path, "rb")
    wb = openpyxl.load_workbook(handle, read_only=True, data_only=True)
    partes: list[str] = []
    total = 0
    try:
        for ws in wb.worksheets:
            partes.append(f"[Planilha: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                celulas = [str(c).strip() for c in row
                           if c is not None and str(c).strip()]
                if not celulas:
                    continue
                linha = " | ".join(celulas)
                partes.append(linha)
                total += len(linha)
                if total > MAX_OCR_CHARS:
                    break
            if total > MAX_OCR_CHARS:
                break
    finally:
        wb.close()
        handle.close()
    return "\n".join(partes)[:MAX_OCR_CHARS]

def _xls_legado(path: str) -> str | None:
    """Texto de .xls legado (xlrd, se instalado). Antes deste ramo, .xls era
    aceito no upload mas indexado VAZIO no RAG — falha silenciosa."""
    try:
        import xlrd  # opcional (requirements)
    except ImportError:
        logger.warning(f"Extração .xls indisponível (xlrd ausente): {path}")
        return None
    linhas: list[str] = []
    book = xlrd.open_workbook(path)
    for sheet in book.sheets():
        linhas.append(f"[Planilha: {sheet.name}]")
        for i in range(sheet.nrows):
            cels = [str(c.value).strip() for c in sheet.row(i) if str(c.value).strip()]
            if cels:
                linhas.append(" | ".join(cels))
            if sum(len(l) for l in linhas) > MAX_OCR_CHARS:
                break
    return "\n".join(linhas)[:MAX_OCR_CHARS] or None


def _doc_legado(path: str) -> str | None:
    """Word binário (.doc): usa antiword/catdoc se disponível no sistema.
    Sem conversor, loga aviso claro — nunca indexa vazio em silêncio."""
    import shutil
    import subprocess
    for conv in ("antiword", "catdoc"):
        exe = shutil.which(conv)
        if exe:
            out = subprocess.run([exe, path], capture_output=True, timeout=60)
            if out.returncode == 0 and out.stdout.strip():
                return out.stdout.decode("utf-8", errors="ignore")[:MAX_OCR_CHARS]
    logger.warning(
        f"Formato .doc sem conversor instalado (antiword/catdoc) — texto não "
        f"extraído; converta para .docx/PDF para indexação RAG: {path}"
    )
    return None



# ── API baseada em BYTES (Fase 3A — /rag/ingest-pdf e tasks Celery) ──────────
# Mesma estratégia da API por caminho acima (texto nativo PyMuPDF + OCR de
# fallback por página), mas recebe o binário direto do UploadFile e retorna
# metadados da extração (nº de páginas, quantas precisaram de OCR).

def ocr_disponivel() -> bool:
    """True se pytesseract E o binário tesseract estão utilizáveis."""
    if not _TESS_OK:
        return False
    import shutil as _shutil
    return _shutil.which("tesseract") is not None


def _ocr_imagem_pil(img) -> str:
    """OCR de um PIL.Image; sem traineddata 'por', degrada para o default."""
    try:
        return pytesseract.image_to_string(img, lang="por")
    except pytesseract.TesseractError:
        return pytesseract.image_to_string(img)


def extrair_texto_pdf(raw: bytes) -> dict:
    """Extrai texto de um PDF em bytes: nativo + OCR nas páginas escaneadas.

    Retorno: {"texto", "paginas", "paginas_ocr", "ocr_disponivel"}.
    Levanta ValueError se o binário não for um PDF legível (chamador → 422).
    Falta de tesseract NUNCA levanta — páginas escaneadas ficam vazias e
    `ocr_disponivel=False` sinaliza a degradação (fallback gracioso).
    """
    if not _PDF_OK:
        raise ValueError("Extração de PDF indisponível no servidor (PyMuPDF ausente)")
    tem_ocr = ocr_disponivel()
    try:
        pdf = fitz.open(stream=raw, filetype="pdf")
    except Exception as e:
        raise ValueError(f"Falha ao ler o PDF: {str(e)[:120]}") from e
    partes: list[str] = []
    paginas_ocr = 0
    with pdf:
        total = pdf.page_count
        for page in pdf:
            txt = page.get_text("text") or ""
            # Página sem texto nativo (escaneada) → OCR, se disponível
            if len(txt.strip()) < 20 and tem_ocr:
                try:
                    pix = page.get_pixmap(dpi=200)
                    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    txt = _ocr_imagem_pil(img)
                    paginas_ocr += 1
                except Exception as e:
                    logger.warning(f"OCR da página {page.number} falhou: {e}")
            partes.append(txt)
            if sum(len(p) for p in partes) > MAX_OCR_CHARS:
                break
    return {
        "texto": "\n".join(partes)[:MAX_OCR_CHARS].strip(),
        "paginas": total,
        "paginas_ocr": paginas_ocr,
        "ocr_disponivel": tem_ocr,
        "por_pagina": [{"pagina": i + 1, "texto": p} for i, p in enumerate(partes)],
    }


def extrair_texto_imagem(raw: bytes) -> dict:
    """OCR de uma imagem (png/jpg...) em bytes. Sem tesseract → texto vazio."""
    import io as _io
    tem_ocr = ocr_disponivel()
    texto = ""
    if tem_ocr:
        try:
            with Image.open(_io.BytesIO(raw)) as img:
                texto = _ocr_imagem_pil(img.convert("RGB"))[:MAX_OCR_CHARS]
        except Exception as e:
            raise ValueError(f"Falha ao ler a imagem: {str(e)[:120]}") from e
    return {
        "texto": texto.strip(),
        "paginas": 1,
        "paginas_ocr": 1 if texto.strip() else 0,
        "ocr_disponivel": tem_ocr,
    }


# ── XML / NF-e ────────────────────────────────────────────────────────────────

def _xml_root(bruto: str):
    """Parse seguro de XML. Sem defusedxml, rejeita DTD/entidades customizadas
    (mitigação de XXE e billion laughs no stdlib)."""
    if not _DEFUSED and re.search(r"<!\s*(DOCTYPE|ENTITY)", bruto, re.I):
        raise ValueError("XML com DTD/ENTITY rejeitado (defusedxml não instalado)")
    return _ET.fromstring(bruto)


def _local(tag) -> str:
    """Nome local do elemento, ignorando namespace ({ns}Tag → Tag)."""
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def _extrair_nfe(root, bruto: str) -> dict | None:
    """
    Se o XML for uma NF-e (nfeProc/NFe/infNFe), extrai campos fiscais:
    chave de acesso (44 dígitos), CNPJ do emitente, valor total (vNF) e NCMs.
    Retorna None se o XML não for NF-e.
    """
    eh_nfe = any(_local(e.tag) in {"nfeProc", "NFe", "infNFe"} for e in root.iter())
    if not eh_nfe:
        return None

    chave = None
    for el in root.iter():
        if _local(el.tag) == "infNFe":
            digits = re.sub(r"\D", "", el.get("Id") or "")
            if len(digits) == 44:
                chave = digits
            break
    if not chave:  # fallback: 44 dígitos consecutivos no texto bruto
        m = re.search(r"(?<!\d)\d{44}(?!\d)", bruto)
        chave = m.group(0) if m else None

    cnpj_emitente = None
    for el in root.iter():
        if _local(el.tag) == "emit":
            for c in el.iter():
                if _local(c.tag) == "CNPJ" and c.text:
                    cnpj_emitente = re.sub(r"\D", "", c.text) or None
                    break
            break

    valor_total = None
    for el in root.iter():
        if _local(el.tag) == "ICMSTot":
            for c in el.iter():
                if _local(c.tag) == "vNF" and c.text:
                    try:
                        valor_total = float(c.text.strip())
                    except ValueError:
                        pass
                    break
            break

    ncms = sorted({
        e.text.strip() for e in root.iter()
        if _local(e.tag) == "NCM" and e.text and e.text.strip()
    })

    return {
        "chave_acesso": chave,
        "cnpj_emitente": cnpj_emitente,
        "valor_total": valor_total,
        "ncms": ncms,
    }


def extrair_xml(filepath: str) -> dict | None:
    """
    Extrai texto pesquisável de um XML e, se for NF-e, campos estruturados.
    Retorno: {"texto": str|None, "nfe": dict|None} ou None em falha de parse.
    Síncrono — chamar via asyncio.to_thread (mesmo padrão de extrair_texto).
    """
    try:
        with open(filepath, encoding="utf-8", errors="ignore") as f:
            bruto = f.read(MAX_OCR_CHARS * 2)
        root = _xml_root(bruto)
    except Exception as e:
        logger.warning(f"Parse XML falhou para {filepath}: {e}")
        return None

    texto = "\n".join(
        t.strip() for t in root.itertext() if t and t.strip()
    )[:MAX_OCR_CHARS] or None

    nfe = None
    try:
        nfe = _extrair_nfe(root, bruto)
    except Exception as e:
        # LGPD: log sem dados do documento — apenas o caminho e o erro
        logger.warning(f"Extração NF-e falhou para {filepath}: {e}")

    return {"texto": texto, "nfe": nfe}
