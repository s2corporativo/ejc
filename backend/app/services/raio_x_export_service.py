"""Exportação profissional do relatório Raio-X para DOCX e PDF."""
from __future__ import annotations

import io
import json
from typing import Any

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fpdf import FPDF


def _text(value: Any) -> str:
    if value is None:
        return "—"
    if isinstance(value, bool):
        return "Sim" if value else "Não"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, indent=2, default=str)
    return str(value)


def _items(report: dict[str, Any], key: str) -> list[Any]:
    value = report.get(key)
    if value is None:
        return []
    return value if isinstance(value, list) else [value]


def gerar_docx(titulo: str, report: dict[str, Any]) -> bytes:
    document = DocxDocument()
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    heading = document.add_heading("RAIO-X DO PROCESSO", level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(titulo)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(report.get("aviso") or "Análise preliminar sujeita à revisão humana.")

    identification = report.get("identificacao") or {}
    document.add_heading("1. Identificação", level=1)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, value in identification.items():
        row = table.add_row().cells
        row[0].text = key.replace("_", " ").title()
        row[1].text = _text(value)

    document.add_heading("2. Síntese executiva", level=1)
    document.add_paragraph(_text(report.get("sintese_executiva")))

    sections = (
        ("3. Partes", "partes"),
        ("4. Cronologia", "cronologia"),
        ("5. Fatos × provas", "fatos_provas"),
        ("6. Pedidos", "pedidos"),
        ("7. Decisões", "decisoes"),
        ("8. Contradições e inconsistências", "contradicoes"),
        ("9. Prazos potenciais", "prazos_potenciais"),
        ("10. Riscos", "riscos"),
        ("11. Pontos fortes", "pontos_fortes"),
        ("12. Pontos frágeis", "pontos_fracos"),
        ("13. Próximos passos", "proximos_passos"),
        ("14. Jornada processual sugerida", "rito_jornada"),
        ("15. Fontes documentais", "fontes"),
    )
    for title, key in sections:
        document.add_heading(title, level=1)
        items = _items(report, key)
        if not items:
            document.add_paragraph("Nenhuma informação identificada com segurança.")
            continue
        for item in items:
            document.add_paragraph(_text(item), style="List Bullet")

    document.add_paragraph()
    document.add_paragraph(
        "Documento gerado pelo EJC como apoio interno. Conferência jurídica humana obrigatória antes de qualquer uso externo.",
    )
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class _Pdf(FPDF):
    def header(self):
        self.set_font("Helvetica", "B", 13)
        self.cell(0, 8, "RAIO-X DO PROCESSO", align="C", new_x="LMARGIN", new_y="NEXT")
        self.ln(2)

    def footer(self):
        self.set_y(-14)
        self.set_font("Helvetica", size=8)
        self.cell(0, 8, f"EJC · página {self.page_no()}", align="C")


def _pdf_text(value: Any) -> str:
    # Fontes core do FPDF não suportam todos os caracteres; substituição segura.
    return _text(value).encode("latin-1", "replace").decode("latin-1")


def gerar_pdf(titulo: str, report: dict[str, Any]) -> bytes:
    pdf = _Pdf()
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 12)
    pdf.multi_cell(0, 7, _pdf_text(titulo), align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=9)
    pdf.multi_cell(
        0,
        5,
        _pdf_text(report.get("aviso") or "Análise preliminar sujeita à revisão humana."),
        new_x="LMARGIN",
        new_y="NEXT",
    )
    pdf.ln(3)

    def section(title: str, value: Any):
        pdf.set_font("Helvetica", "B", 11)
        pdf.multi_cell(0, 7, _pdf_text(title), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", size=9)
        values = value if isinstance(value, list) else [value]
        values = [item for item in values if item not in (None, "", [], {})]
        if not values:
            pdf.multi_cell(0, 5, "Nenhuma informacao identificada com seguranca.", new_x="LMARGIN", new_y="NEXT")
        for item in values:
            pdf.multi_cell(0, 5, _pdf_text(item), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        pdf.ln(2)

    section("1. Identificação", report.get("identificacao"))
    section("2. Síntese executiva", report.get("sintese_executiva"))
    section("3. Partes", report.get("partes"))
    section("4. Cronologia", report.get("cronologia"))
    section("5. Fatos e provas", report.get("fatos_provas"))
    section("6. Pedidos", report.get("pedidos"))
    section("7. Decisões", report.get("decisoes"))
    section("8. Contradições", report.get("contradicoes"))
    section("9. Prazos potenciais", report.get("prazos_potenciais"))
    section("10. Riscos", report.get("avaliacao_risco") or report.get("riscos"))
    section("11. Próximos passos", report.get("proximos_passos"))
    section("12. Jornada sugerida", report.get("rito_jornada"))
    section("13. Fontes", report.get("fontes"))
    return bytes(pdf.output())
